from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "report-source.md"
OUTPUT = ROOT / "OLP-2026-DX-OS-project-recommendations.docx"

INK = "111827"
MUTED = "52606D"
TEAL = "0F766E"
TEAL_DARK = "115E59"
TEAL_SOFT = "E8F4F2"
BLUE_SOFT = "EEF4FA"
LINE = "CBD5E1"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL_DARK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(color)
    run_pr.append(underline)
    text = OxmlElement("w:t")
    text.text = label
    run.append(run_pr)
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^\s)]+(?:\)[^\s)]*)?\))")


def add_inline(paragraph, text: str) -> None:
    pos = 0
    # A deliberately small Markdown inline parser; source avoids nested markup.
    token_re = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^\s]+?\)(?=$|[\s.,;:]))")
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(TEAL_DARK)
        else:
            label, url = re.match(r"\[([^\]]+)\]\((https?://.+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Title", 26, TEAL_DARK, 0, 12),
        ("Heading 1", 17, INK, 14, 6),
        ("Heading 2", 12.5, TEAL_DARK, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Executive" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Executive", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12

    if "Caption Small" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Caption Small", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Aptos"
        style.font.size = Pt(8)
        style.font.color.rgb = RGBColor.from_string(MUTED)
        style.paragraph_format.space_after = Pt(3)

    props = doc.core_properties
    props.title = "Đề xuất chủ đề dự án OLP PMNM 2026 theo mô hình DX-OS"
    props.subject = "Nghiên cứu và khuyến nghị chiến lược cho vòng loại HUTECH và OLP 2026"
    props.author = "Nhóm dự án OLP 2026"
    props.keywords = "DX-OS, OLP 2026, phần mềm nguồn mở, AI agent, doanh nghiệp"


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("OLP PMNM 2026  /  DX-OS")
    r.font.name = "Aptos Mono"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.add_run("Đề xuất chủ đề dự án\n")
    accent = title.add_run("Hệ điều hành doanh nghiệp số AI")
    accent.font.color.rgb = RGBColor.from_string(TEAL)
    title.paragraph_format.space_after = Pt(16)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(22)
    run = line.add_run("OPEN DX  ×  ENTERPRISE AI GOVERNANCE")
    run.font.name = "Aptos Mono"
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)

    thesis = doc.add_paragraph()
    thesis.paragraph_format.space_before = Pt(10)
    thesis.paragraph_format.space_after = Pt(16)
    thesis.paragraph_format.left_indent = Inches(0.28)
    thesis.paragraph_format.right_indent = Inches(0.55)
    thesis.paragraph_format.line_spacing = 1.12
    thesis._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))
    r = thesis.add_run("Khuyến nghị số 1\n")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r = thesis.add_run("OpenDX Agent Passport")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(INK)
    r = thesis.add_run("\nHộ chiếu nhiệm vụ và cổng kiểm soát AI Agent cho doanh nghiệp")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_paragraph("Bản nghiên cứu định hướng cho vòng loại HUTECH và đội tuyển OLP quốc gia", style="Executive")
    doc.add_paragraph("Ngày chốt thông tin: 07/09/2026  •  Địa điểm OLP quốc gia: VKU, Đà Nẵng", style="Caption Small")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(110)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Tóm tắt quyết định")
    r.font.name = "Aptos Mono"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL_DARK)
    p = doc.add_paragraph(style="Executive")
    p.paragraph_format.right_indent = Inches(0.5)
    add_inline(p, "Không làm thêm một chatbot. Xây một lớp vận hành chứng minh **ai giao việc, agent được phép làm gì, dữ liệu nào được dùng, hành động nào phải duyệt và toàn bộ việc đó được truy vết ra sao**.")

    doc.add_page_break()


def setup_headers_footers(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.text = "OLP PMNM 2026  /  DX-OS PROJECT RESEARCH"
    p.style = doc.styles["Caption Small"]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    r = p.add_run("HUTECH → OLP quốc gia   •   07/09/2026")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    p.add_run("                                                   ")
    add_page_number(p)


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        rows.append([cell.strip() for cell in raw.split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    available = 7.05
    max_lens = []
    for c in range(cols):
        values = [len(re.sub(r"\[[^\]]+\]\([^)]+\)", "link", r[c] if c < len(r) else "")) for r in rows]
        max_lens.append(max(7, min(max(values), 42)))
    if cols >= 7:
        widths = [0.36, 2.15] + [(available - 2.51) / (cols - 2)] * (cols - 2)
        font_size = 6.8
    else:
        total = sum(max_lens)
        widths = [available * x / total for x in max_lens]
        widths = [max(0.62, w) for w in widths]
        scale = available / sum(widths)
        widths = [w * scale for w in widths]
        font_size = 7.8 if cols >= 4 else 8.2

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        prevent_row_split(row)
        for c in range(cols):
            cell = row.cells[c]
            cell.width = Inches(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 70, 75, 70, 75)
            if r_idx == 0:
                set_cell_shading(cell, TEAL_DARK)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, TEAL_SOFT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if c >= 2 and cols >= 7:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = row_data[c] if c < len(row_data) else ""
            add_inline(p, text)
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(font_size)
                if r_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body_from_markdown(doc: Document, source: str) -> None:
    lines = source.splitlines()
    # Skip title and cover metadata; body begins at first level-2 heading.
    i = next(idx for idx, line in enumerate(lines) if line.startswith("## "))
    paragraph_buffer: list[str] = []
    exec_summary = True

    def flush() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(s.strip().removesuffix("  ") for s in paragraph_buffer).strip()
        style = "Executive" if exec_summary else None
        p = doc.add_paragraph(style=style)
        add_inline(p, text)
        paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("### "):
            flush()
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("## "):
            flush()
            title = line[3:].strip()
            if title.startswith("5."):
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, title)
            # A thin rule below major headings.
            p_pr = p._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "10")
            bottom.set(qn("w:space"), "4")
            bottom.set(qn("w:color"), TEAL)
            p_bdr.append(bottom)
            p_pr.append(p_bdr)
            exec_summary = title == "Kết luận điều hành"
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)", stripped)
        number = re.match(r"^(\d+)\.\s+(.+)", stripped)
        if bullet or number:
            flush()
            p = doc.add_paragraph(style="List Bullet" if bullet else None)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            if number:
                p.paragraph_format.left_indent = Inches(0.24)
                p.paragraph_format.first_line_indent = Inches(-0.24)
                marker = p.add_run(f"{number.group(1)}.  ")
                marker.bold = True
                marker.font.color.rgb = RGBColor.from_string(MUTED)
                add_inline(p, number.group(2))
            else:
                add_inline(p, bullet.group(1))
            i += 1
            continue
        paragraph_buffer.append(line)
        i += 1
    flush()


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    setup_headers_footers(doc)
    add_body_from_markdown(doc, source)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
